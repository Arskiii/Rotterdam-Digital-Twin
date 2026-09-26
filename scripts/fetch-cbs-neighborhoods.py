#!/usr/bin/env python3
"""Extract public, aggregated Rotterdam neighbourhood figures from CBS KWB 2025.

Requires openpyxl. The original workbook is downloaded from CBS and is not
redistributed; the small output carries only neighborhood totals.
"""
import json
import math
import re
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
URL = 'https://download.cbs.nl/regionale-kaarten/kwb2025.xlsx'
TARGET = ROOT / 'public' / 'data' / 'cbs-neighborhoods.json'
WORKBOOK = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('/tmp/kwb2025.xlsx')
if len(sys.argv) == 1:
    request = urllib.request.Request(URL, headers={'User-Agent': 'Rotterdam-Digital-Twin/1.0 (public CBS statistics)'})
    with urllib.request.urlopen(request, timeout=90) as response, WORKBOOK.open('wb') as output:
        total = 0
        while chunk := response.read(1024 * 1024):
            total += len(chunk)
            if total > 100 * 1024 * 1024:
                raise SystemExit('CBS workbook exceeds 100 MB safety limit')
            output.write(chunk)
wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
required = {'gwb_code', 'a_inw', 'a_65_oo'}
rows = []
for sheet in wb:
    iterator = sheet.iter_rows(values_only=True)
    header = None
    for row in iterator:
        cells = [str(value).strip().lower() if value is not None else '' for value in row]
        if required.issubset(cells):
            header = {value: i for i, value in enumerate(cells) if value}
            break
    if not header:
        continue
    for row in iterator:
        code = str(row[header['gwb_code']] or '').strip().upper()
        if not re.fullmatch(r'BU0599[A-Z0-9]{4}', code):
            continue
        def number(key):
            if key not in header: return None
            value = row[header[key]]
            if value is None or str(value).strip() in ('', '.', 'x', '-'): return None
            try:
                result = float(str(value).replace(',', '.'))
            except ValueError:
                return None
            return result if math.isfinite(result) else None
        population = number('a_inw')
        if population is None or population < 0: continue
        name = str(row[header['regio']]) if 'regio' in header else code
        rows.append({
            'code': code, 'name': name.strip(), 'residents': round(population),
            'age65Plus': round(number('a_65_oo')) if number('a_65_oo') is not None else None,
            'households': round(number('a_hh')) if number('a_hh') is not None else None,
            'lowIncomeHouseholdsPct': number('p_hh_lkk'),
        })
    if rows: break
if len(rows) < 50:
    raise SystemExit(f'CBS parse failed: expected at least 50 Rotterdam neighbourhoods, found {len(rows)}')
rows.sort(key=lambda row: row['code'])
TARGET.write_text(json.dumps({
    'schemaVersion': 1, 'publisher': 'CBS', 'dataset': 'Kerncijfers wijken en buurten 2025',
    'datasetId': '86165NED', 'source': URL, 'retrievedAt': datetime.now(timezone.utc).isoformat(),
    'license': 'CC BY 4.0', 'geography': 'CBS BU0599 neighbourhoods, municipality Rotterdam',
    'caveat': 'Aggregated released statistics; suppressed or missing values are null; geography is not the simulation district boundary.',
    'neighborhoods': rows,
}, ensure_ascii=False, separators=(',', ':')) + '\n')
print(f'Wrote {len(rows)} Rotterdam neighbourhoods to {TARGET}')
